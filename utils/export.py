# -*- coding: utf-8 -*-
"""Export des résultats en Word et Excel avec graphiques."""

import os
import tempfile
import pandas as pd
import plotly.io as pio
from docx import Document
from docx.shared import Pt, Cm, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT


# ==============================================================================
# FONCTIONS INTERNES WORD
# ==============================================================================
def _ajouter_titre(doc, texte, niveau=1):
    """Ajoute un titre au document Word."""
    h = doc.add_heading(texte, level=niveau)
    for run in h.runs:
        run.font.color.rgb = RGBColor(0x1F, 0x3B, 0x73)
    return h


def _ajouter_tableau(doc, df, max_lignes=None):
    """Ajoute un DataFrame sous forme de tableau Word."""
    if max_lignes:
        df = df.head(max_lignes)
    table = doc.add_table(rows=1, cols=len(df.columns))
    table.style = "Light Grid Accent 1"
    table.alignment = WD_TABLE_ALIGNMENT.CENTER

    for i, col in enumerate(df.columns):
        cellule = table.rows[0].cells[i]
        cellule.text = str(col)
        for p in cellule.paragraphs:
            for r in p.runs:
                r.bold = True
                r.font.size = Pt(10)

    for _, ligne in df.iterrows():
        cells = table.add_row().cells
        for i, val in enumerate(ligne):
            cells[i].text = str(val)
            for p in cells[i].paragraphs:
                for r in p.runs:
                    r.font.size = Pt(9)
    return table


def _ajouter_graphique_word(doc, figure_plotly, titre=None):
    """Convertit un graphique Plotly en PNG et l'insère dans Word."""
    try:
        img_bytes = pio.to_image(
            figure_plotly, format="png",
            width=900, height=500, scale=2,
        )
        with tempfile.NamedTemporaryFile(delete=False, suffix=".png") as tmp:
            tmp.write(img_bytes)
            chemin = tmp.name
        if titre:
            p = doc.add_paragraph()
            run = p.add_run(titre)
            run.bold = True
        doc.add_picture(chemin, width=Cm(16))
        os.unlink(chemin)
    except Exception as e:
        doc.add_paragraph(f"[Graphique non disponible : {e}]")


# ==============================================================================
# EXPORT WORD COMPLET
# ==============================================================================
def exporter_word_complet(chemin_sortie, titre_rapport, sections):
    """
    sections : liste de dict avec clés :
        - titre (str)
        - contenu (str, optionnel)
        - tableau (DataFrame, optionnel)
        - graphique (Figure Plotly, optionnel)
    """
    doc = Document()
    style = doc.styles["Normal"]
    style.font.name = "Calibri"
    style.font.size = Pt(11)

    # Page de titre
    h = doc.add_heading(titre_rapport, level=0)
    h.alignment = WD_ALIGN_PARAGRAPH.CENTER
    doc.add_paragraph()
    doc.add_page_break()

    for section in sections:
        _ajouter_titre(doc, section["titre"], niveau=1)

        if section.get("contenu"):
            doc.add_paragraph(section["contenu"])

        if section.get("tableau") is not None:
            _ajouter_tableau(doc, section["tableau"])

        if section.get("graphique") is not None:
            _ajouter_graphique_word(doc, section["graphique"])

        doc.add_paragraph()

    doc.save(chemin_sortie)
    return chemin_sortie


# ==============================================================================
# EXPORT EXCEL ENRICHI
# ==============================================================================
def exporter_excel_enrichi(dictionnaires, chemin_sortie):
    """
    Export multi-feuilles avec styles : titres en gras, entêtes en bleu,
    colonnes auto-dimensionnées.
    dictionnaires : dict {nom_feuille: DataFrame}
    """
    with pd.ExcelWriter(chemin_sortie, engine="xlsxwriter") as writer:
        workbook = writer.book

        # Formats réutilisables
        fmt_titre = workbook.add_format({
            "bold": True,
            "font_size": 14,
            "font_color": "#1F3B73",
            "align": "center",
            "valign": "vcenter",
        })
        fmt_entete = workbook.add_format({
            "bold": True,
            "bg_color": "#1F3B73",
            "font_color": "white",
            "border": 1,
            "align": "center",
            "valign": "vcenter",
        })
        fmt_cellule = workbook.add_format({
            "border": 1,
            "align": "left",
            "valign": "vcenter",
        })
        fmt_nombre = workbook.add_format({
            "border": 1,
            "num_format": "0.00",
            "align": "center",
        })

        for nom_feuille, df in dictionnaires.items():
            nom_court = nom_feuille[:31]

            # Sécurité : si le DataFrame est vide, on met un placeholder
            if df is None or (isinstance(df, pd.DataFrame) and df.empty and len(df.columns) == 0):
                pd.DataFrame({"Message": ["Aucune donnée"]}).to_excel(
                    writer, sheet_name=nom_court, index=False, startrow=1,
                )
                worksheet = writer.sheets[nom_court]
                worksheet.merge_range(0, 0, 0, 0, nom_feuille, fmt_titre)
                continue

            df.to_excel(writer, sheet_name=nom_court, index=False, startrow=1)
            worksheet = writer.sheets[nom_court]

            # Titre fusionné
            worksheet.merge_range(
                0, 0, 0, max(len(df.columns) - 1, 0),
                nom_feuille, fmt_titre,
            )

            # Entêtes
            for i, col in enumerate(df.columns):
                worksheet.write(1, i, col, fmt_entete)

            # Largeur et format des colonnes
            for i, col in enumerate(df.columns):
                try:
                    longueur_max = df[col].astype(str).str.len().max()
                    if pd.isna(longueur_max):
                        longueur_max = 10
                except Exception:
                    longueur_max = 10
                largeur = max(len(str(col)), int(longueur_max)) + 2

                if pd.api.types.is_numeric_dtype(df[col]):
                    worksheet.set_column(i, i, min(largeur, 15), fmt_nombre)
                else:
                    worksheet.set_column(i, i, min(largeur, 40), fmt_cellule)

            # Figer la ligne des entêtes
            worksheet.freeze_panes(2, 0)

    return chemin_sortie


# ==============================================================================
# ALIAS DE COMPATIBILITÉ (au cas où un ancien code appelle exporter_excel)
# ==============================================================================
def exporter_excel(dictionnaires, chemin_sortie):
    """Alias vers exporter_excel_enrichi."""
    return exporter_excel_enrichi(dictionnaires, chemin_sortie)